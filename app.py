import csv
import io
import os
import sqlite3
from datetime import datetime, date, timedelta
from flask import Flask, render_template, request, redirect, url_for, flash, Response, session

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or os.urandom(24)

CRM_PASSWORD = os.environ.get("CRM_PASSWORD")
if not CRM_PASSWORD:
    raise RuntimeError("CRM_PASSWORD environment variable is not set.")


@app.before_request
def require_login():
    if request.endpoint in ("login", "logout", "static"):
        return
    if not session.get("logged_in"):
        return redirect(url_for("login", next=request.path))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if request.form.get("password") == CRM_PASSWORD:
            session["logged_in"] = True
            session.permanent = True
            next_url = request.form.get("next") or url_for("index")
            return redirect(next_url)
        flash("Wrong password.", "danger")
    return render_template("login.html", next=request.args.get("next", ""))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.template_filter("fmtdatetime")
def fmt_datetime(value):
    """Format 'YYYY-MM-DD HH:MM' or 'YYYY-MM-DD' into 'Feb 20 · 2:30 PM'."""
    if not value:
        return ""
    s = str(value)
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            dt = datetime.strptime(s, fmt)
            date_str = dt.strftime("%b") + " " + str(dt.day)
            if " " in s:
                hour = dt.hour % 12 or 12
                ampm = "AM" if dt.hour < 12 else "PM"
                time_str = f"{hour}:{dt.minute:02d} {ampm}"
                return f"{date_str} · {time_str}"
            return date_str
        except ValueError:
            continue
    return s

DATABASE = os.environ.get("DATABASE_PATH", "crm.db")


def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS deals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            asset_type TEXT,
            location TEXT,
            target_raise TEXT,
            status TEXT DEFAULT 'Active',
            notes TEXT,
            purchase_price TEXT,
            units_sf TEXT,
            noi_current TEXT,
            noi_projected TEXT,
            cap_rate TEXT,
            exit_cap_rate TEXT,
            ltv TEXT,
            hold_period TEXT,
            irr_target TEXT,
            equity_multiple TEXT,
            closing_date TEXT,
            email_keywords TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS outreach (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            deal_id INTEGER NOT NULL,
            contact_name TEXT NOT NULL,
            company TEXT,
            email TEXT,
            sent_by TEXT,
            outreach_date TEXT,
            outreach_summary TEXT,
            response_date TEXT,
            response_summary TEXT,
            follow_up_needed INTEGER DEFAULT 0,
            follow_up_date TEXT,
            stage TEXT DEFAULT 'Initial Email',
            interest_level TEXT DEFAULT 'Unknown',
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (deal_id) REFERENCES deals(id)
        );

        CREATE TABLE IF NOT EXISTS gp_contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            title TEXT,
            company TEXT,
            email TEXT,
            phone TEXT,
            location TEXT,
            aum_range TEXT,
            strategy TEXT,
            min_check TEXT,
            max_check TEXT,
            source TEXT,
            status TEXT DEFAULT 'New',
            notes TEXT,
            last_contacted TEXT,
            next_contact_date TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS contact_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            contact_id INTEGER NOT NULL,
            note_date TEXT,
            note_text TEXT NOT NULL,
            note_html TEXT,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (contact_id) REFERENCES gp_contacts(id)
        );

        CREATE TABLE IF NOT EXISTS lp_investor_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            investor_id INTEGER NOT NULL,
            note_date TEXT,
            note_text TEXT NOT NULL,
            note_html TEXT,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (investor_id) REFERENCES lp_investors(id)
        );

        CREATE TABLE IF NOT EXISTS email_templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            subject TEXT,
            body TEXT NOT NULL,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS emails (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject TEXT,
            body TEXT,
            from_addr TEXT,
            to_addr TEXT,
            cc_addr TEXT,
            sent_on TEXT,
            direction TEXT DEFAULT 'sent',
            tags TEXT,
            deal_id INTEGER,
            contact_id INTEGER,
            entry_id TEXT UNIQUE,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (deal_id) REFERENCES deals(id),
            FOREIGN KEY (contact_id) REFERENCES gp_contacts(id)
        );

        CREATE TABLE IF NOT EXISTS lp_investors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            title TEXT,
            company TEXT,
            email TEXT,
            phone TEXT,
            location TEXT,
            net_worth TEXT,
            preferred_asset_types TEXT,
            preferred_markets TEXT,
            min_check TEXT,
            max_check TEXT,
            hold_period_pref TEXT,
            accredited TEXT DEFAULT 'Yes',
            source TEXT,
            status TEXT DEFAULT 'New',
            notes TEXT,
            last_contacted TEXT,
            next_contact_date TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS lenders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            title TEXT,
            company TEXT,
            email TEXT,
            phone TEXT,
            lender_type TEXT,
            loan_types TEXT,
            asset_types TEXT,
            min_loan TEXT,
            max_loan TEXT,
            max_ltv TEXT,
            typical_rate TEXT,
            recourse TEXT,
            markets TEXT,
            typical_term TEXT,
            origination_fee TEXT,
            status TEXT DEFAULT 'New',
            notes TEXT,
            last_contacted TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS import_batches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            label TEXT NOT NULL,
            type TEXT NOT NULL,
            count INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS campaigns (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            subject TEXT NOT NULL,
            body_template TEXT NOT NULL,
            status TEXT DEFAULT 'draft',
            total_count INTEGER DEFAULT 0,
            sent_count INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now')),
            sent_at TEXT
        );

        CREATE TABLE IF NOT EXISTS campaign_emails (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            campaign_id INTEGER NOT NULL,
            contact_id INTEGER,
            to_name TEXT,
            to_email TEXT,
            company TEXT,
            subject TEXT,
            body TEXT,
            status TEXT DEFAULT 'pending',
            created_at TEXT DEFAULT (datetime('now')),
            sent_at TEXT,
            FOREIGN KEY (campaign_id) REFERENCES campaigns(id),
            FOREIGN KEY (contact_id) REFERENCES gp_contacts(id)
        );
    """)

    # Safe migrations for existing databases
    for stmt in [
        "ALTER TABLE outreach ADD COLUMN follow_up_date TEXT",
        "ALTER TABLE gp_contacts ADD COLUMN next_contact_date TEXT",
        "ALTER TABLE outreach ADD COLUMN stage TEXT DEFAULT 'Initial Email'",
        "ALTER TABLE outreach ADD COLUMN interest_level TEXT DEFAULT 'Unknown'",
        "ALTER TABLE deals ADD COLUMN purchase_price TEXT",
        "ALTER TABLE deals ADD COLUMN units_sf TEXT",
        "ALTER TABLE deals ADD COLUMN noi_current TEXT",
        "ALTER TABLE deals ADD COLUMN noi_projected TEXT",
        "ALTER TABLE deals ADD COLUMN cap_rate TEXT",
        "ALTER TABLE deals ADD COLUMN exit_cap_rate TEXT",
        "ALTER TABLE deals ADD COLUMN ltv TEXT",
        "ALTER TABLE deals ADD COLUMN hold_period TEXT",
        "ALTER TABLE deals ADD COLUMN irr_target TEXT",
        "ALTER TABLE deals ADD COLUMN equity_multiple TEXT",
        "ALTER TABLE deals ADD COLUMN closing_date TEXT",
        "ALTER TABLE emails ADD COLUMN from_addr TEXT",
        "ALTER TABLE emails ADD COLUMN direction TEXT DEFAULT 'sent'",
        "ALTER TABLE emails ADD COLUMN tags TEXT",
        "ALTER TABLE deals ADD COLUMN email_keywords TEXT",
        "ALTER TABLE campaigns ADD COLUMN target_audience TEXT DEFAULT 'gp'",
        "ALTER TABLE campaign_emails ADD COLUMN lp_investor_id INTEGER",
        "ALTER TABLE gp_contacts ADD COLUMN import_batch_id INTEGER",
        "ALTER TABLE lp_investors ADD COLUMN import_batch_id INTEGER",
        "ALTER TABLE contact_notes ADD COLUMN note_html TEXT",
    ]:
        try:
            conn.execute(stmt)
        except Exception:
            pass

    # Seed default email templates once
    if conn.execute("SELECT COUNT(*) FROM email_templates").fetchone()[0] == 0:
        seed = [
            (
                "Initial Deal Introduction",
                "Introduction: [DEAL NAME]",
                "Hi [CONTACT NAME],\n\nI hope this finds you well. I wanted to reach out regarding an opportunity we're currently marketing that I thought might be a strong fit for your investment criteria.\n\n[DEAL NAME] is a [ASSET TYPE] opportunity located in [LOCATION], with a target equity raise of [TARGET RAISE]. A few highlights:\n\n\u2022 [HIGHLIGHT 1]\n\u2022 [HIGHLIGHT 2]\n\u2022 [HIGHLIGHT 3]\n\nHappy to send over the full investment summary if you have interest. Would you have 15 minutes for a quick call this week?\n\nBest,\n[YOUR NAME]",
            ),
            (
                "Follow-Up (No Response)",
                "Following Up: [DEAL NAME]",
                "Hi [CONTACT NAME],\n\nJust wanted to follow up on my previous note regarding [DEAL NAME] \u2014 I know things get busy and wanted to make sure this didn\u2019t fall through the cracks.\n\nHappy to send the investment summary or get on a quick call at your convenience.\n\nBest,\n[YOUR NAME]",
            ),
            (
                "Follow-Up After Interest",
                "Next Steps: [DEAL NAME]",
                "Hi [CONTACT NAME],\n\nGreat speaking with you. As discussed, I\u2019m sending over the materials for [DEAL NAME].\n\nPlease don\u2019t hesitate to reach out with any questions. I\u2019d suggest we reconnect in two weeks to discuss further \u2014 does that work on your end?\n\nBest,\n[YOUR NAME]",
            ),
            (
                "Meeting Request",
                "Meeting Request: [DEAL NAME]",
                "Hi [CONTACT NAME],\n\nThank you for your interest in [DEAL NAME]. I\u2019d love to set up a call to walk you through the deal in more detail.\n\nAre any of the following times convenient?\n\n\u2022 [TIME OPTION 1]\n\u2022 [TIME OPTION 2]\n\u2022 [TIME OPTION 3]\n\nOr feel free to send over a time that works better for you.\n\nBest,\n[YOUR NAME]",
            ),
        ]
        conn.executemany(
            "INSERT INTO email_templates (name, subject, body) VALUES (?, ?, ?)", seed
        )

    conn.commit()
    conn.close()


# ── Helpers ────────────────────────────────────────────────────────────────────

def parse_money(s):
    """Parse '$25M', '$500K', '$1B\u2013$2B' etc. into a float."""
    if not s:
        return None
    s = s.replace("$", "").replace(",", "").replace(" ", "").upper()
    for sep in ["\u2013", "\u2014", "-", "TO"]:
        if sep in s:
            parts = s.split(sep, 1)
            vals = [parse_money(p) for p in parts]
            vals = [v for v in vals if v is not None]
            return sum(vals) / len(vals) if vals else None
    mult = 1
    if s.endswith("B"):
        mult = 1_000_000_000; s = s[:-1]
    elif s.endswith("M"):
        mult = 1_000_000; s = s[:-1]
    elif s.endswith("K"):
        mult = 1_000; s = s[:-1]
    try:
        return float(s) * mult
    except Exception:
        return None


def score_gp(deal, gp):
    """Return (int score, [str reasons]) for a GP against a deal."""
    score = 0
    reasons = []
    deal_raise = parse_money(deal["target_raise"])
    gp_min = parse_money(gp["min_check"])
    gp_max = parse_money(gp["max_check"])

    if deal_raise and gp_min and gp_max:
        if gp_min <= deal_raise <= gp_max:
            score += 3; reasons.append("check size fits")
        elif deal_raise <= gp_max:
            score += 1; reasons.append("within max check")
    elif deal_raise and gp_max and deal_raise <= gp_max:
        score += 1; reasons.append("within max check")

    # Strategy keyword overlap
    if deal["asset_type"] and gp["strategy"]:
        stop = {"and", "or", "the", "in", "of", "for", "a", "an", "with",
                "value", "add", "core", "plus", "class", "real", "estate"}
        dw = set(deal["asset_type"].lower().split()) - stop
        gw = set(gp["strategy"].lower().split()) - stop
        if dw & gw:
            score += 2; reasons.append("strategy match")

    # Location overlap (city-level)
    if deal["location"] and gp["location"]:
        city = deal["location"].lower().split(",")[0].strip()
        if city and city in gp["location"].lower():
            score += 1; reasons.append("same market")

    return score, reasons


# ── Template context ───────────────────────────────────────────────────────────

@app.context_processor
def inject_globals():
    conn = get_db()
    gp_contacts = conn.execute("SELECT id, name, company FROM gp_contacts ORDER BY name").fetchall()
    deals = conn.execute("SELECT id, name FROM deals ORDER BY name").fetchall()
    conn.close()
    return dict(_all_gp_contacts=gp_contacts, _all_deals=deals, today=date.today().isoformat())


# ── Dashboard ──────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    today = date.today().isoformat()
    conn = get_db()

    deals = conn.execute(
        "SELECT d.*, COUNT(o.id) as outreach_count "
        "FROM deals d LEFT JOIN outreach o ON d.id = o.deal_id "
        "GROUP BY d.id ORDER BY d.created_at DESC"
    ).fetchall()

    gp_count = conn.execute("SELECT COUNT(*) FROM gp_contacts").fetchone()[0]
    pending_followups = conn.execute(
        "SELECT COUNT(*) FROM outreach WHERE follow_up_needed = 1"
    ).fetchone()[0]

    overdue_outreach = conn.execute(
        "SELECT o.*, d.name as deal_name FROM outreach o "
        "JOIN deals d ON o.deal_id = d.id "
        "WHERE o.follow_up_needed = 1 AND o.follow_up_date IS NOT NULL AND o.follow_up_date <= ? "
        "ORDER BY o.follow_up_date ASC",
        (today,),
    ).fetchall()

    overdue_contacts = conn.execute(
        "SELECT * FROM gp_contacts "
        "WHERE next_contact_date IS NOT NULL AND next_contact_date <= ? "
        "ORDER BY next_contact_date ASC",
        (today,),
    ).fetchall()

    recent_activity = conn.execute(
        "SELECT o.*, d.name as deal_name FROM outreach o "
        "JOIN deals d ON o.deal_id = d.id "
        "ORDER BY COALESCE(o.outreach_date, o.created_at) DESC LIMIT 8"
    ).fetchall()

    conn.close()
    return render_template(
        "index.html",
        deals=deals,
        gp_count=gp_count,
        pending_followups=pending_followups,
        overdue_outreach=overdue_outreach,
        overdue_contacts=overdue_contacts,
        recent_activity=recent_activity,
        today=today,
    )


# ── Deals ──────────────────────────────────────────────────────────────────────

@app.route("/deals")
def deals():
    conn = get_db()
    rows = conn.execute(
        "SELECT d.*, COUNT(o.id) as outreach_count "
        "FROM deals d LEFT JOIN outreach o ON d.id = o.deal_id "
        "GROUP BY d.id ORDER BY d.created_at DESC"
    ).fetchall()
    conn.close()
    return render_template("deals.html", deals=rows)


@app.route("/deals/new", methods=["GET", "POST"])
def new_deal():
    if request.method == "POST":
        conn = get_db()
        conn.execute(
            "INSERT INTO deals (name, asset_type, location, target_raise, status, notes, "
            "purchase_price, units_sf, noi_current, noi_projected, cap_rate, exit_cap_rate, "
            "ltv, hold_period, irr_target, equity_multiple, closing_date, email_keywords) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (request.form["name"], request.form.get("asset_type", ""),
             request.form.get("location", ""), request.form.get("target_raise", ""),
             request.form.get("status", "Active"), request.form.get("notes", ""),
             request.form.get("purchase_price", "") or None, request.form.get("units_sf", "") or None,
             request.form.get("noi_current", "") or None, request.form.get("noi_projected", "") or None,
             request.form.get("cap_rate", "") or None, request.form.get("exit_cap_rate", "") or None,
             request.form.get("ltv", "") or None, request.form.get("hold_period", "") or None,
             request.form.get("irr_target", "") or None, request.form.get("equity_multiple", "") or None,
             request.form.get("closing_date", "") or None,
             request.form.get("email_keywords", "") or None),
        )
        conn.commit(); conn.close()
        flash("Deal created.", "success")
        return redirect(url_for("deals"))
    return render_template("deal_form.html", deal=None, title="New Deal")


@app.route("/deals/<int:deal_id>")
def deal_detail(deal_id):
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger"); return redirect(url_for("deals"))
    outreaches = conn.execute(
        "SELECT * FROM outreach WHERE deal_id = ? ORDER BY outreach_date DESC, created_at DESC",
        (deal_id,),
    ).fetchall()
    all_gps = conn.execute("SELECT * FROM gp_contacts").fetchall()
    matches = []
    for gp in all_gps:
        sc, reasons = score_gp(deal, gp)
        if sc > 0:
            matches.append({"gp": gp, "score": sc, "reasons": reasons})
    matches.sort(key=lambda x: x["score"], reverse=True)
    conn.close()
    return render_template("deal_detail.html", deal=deal, outreaches=outreaches, matching_gps=matches[:10])


@app.route("/deals/<int:deal_id>/edit", methods=["GET", "POST"])
def edit_deal(deal_id):
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger"); return redirect(url_for("deals"))
    if request.method == "POST":
        conn.execute(
            "UPDATE deals SET name=?, asset_type=?, location=?, target_raise=?, status=?, notes=?, "
            "purchase_price=?, units_sf=?, noi_current=?, noi_projected=?, cap_rate=?, exit_cap_rate=?, "
            "ltv=?, hold_period=?, irr_target=?, equity_multiple=?, closing_date=?, email_keywords=? WHERE id=?",
            (request.form["name"], request.form.get("asset_type", ""),
             request.form.get("location", ""), request.form.get("target_raise", ""),
             request.form.get("status", "Active"), request.form.get("notes", ""),
             request.form.get("purchase_price", "") or None, request.form.get("units_sf", "") or None,
             request.form.get("noi_current", "") or None, request.form.get("noi_projected", "") or None,
             request.form.get("cap_rate", "") or None, request.form.get("exit_cap_rate", "") or None,
             request.form.get("ltv", "") or None, request.form.get("hold_period", "") or None,
             request.form.get("irr_target", "") or None, request.form.get("equity_multiple", "") or None,
             request.form.get("closing_date", "") or None,
             request.form.get("email_keywords", "") or None, deal_id),
        )
        conn.commit(); conn.close()
        flash("Deal updated.", "success")
        return redirect(url_for("deal_detail", deal_id=deal_id))
    conn.close()
    return render_template("deal_form.html", deal=deal, title="Edit Deal")


@app.route("/deals/<int:deal_id>/delete", methods=["POST"])
def delete_deal(deal_id):
    conn = get_db()
    conn.execute("DELETE FROM outreach WHERE deal_id = ?", (deal_id,))
    conn.execute("DELETE FROM deals WHERE id = ?", (deal_id,))
    conn.commit(); conn.close()
    flash("Deal deleted.", "info")
    return redirect(url_for("deals"))


@app.route("/deals/export.csv")
def export_deals():
    conn = get_db()
    rows = conn.execute("SELECT * FROM deals ORDER BY created_at DESC").fetchall()
    conn.close()
    cols = ["id", "name", "asset_type", "location", "target_raise", "status",
            "purchase_price", "units_sf", "noi_current", "noi_projected",
            "cap_rate", "exit_cap_rate", "ltv", "hold_period",
            "irr_target", "equity_multiple", "closing_date", "notes", "created_at"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cols)
    for r in rows:
        writer.writerow([r[c] for c in cols])
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=deals.csv"})


@app.route("/deals/<int:deal_id>/outreach/export.csv")
def export_outreach(deal_id):
    conn = get_db()
    deal = conn.execute("SELECT name FROM deals WHERE id = ?", (deal_id,)).fetchone()
    rows = conn.execute(
        "SELECT * FROM outreach WHERE deal_id = ? ORDER BY outreach_date DESC", (deal_id,)
    ).fetchall()
    conn.close()
    cols = ["id", "contact_name", "company", "email", "sent_by", "outreach_date",
            "outreach_summary", "response_date", "response_summary",
            "follow_up_needed", "stage", "interest_level", "created_at"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cols)
    for r in rows:
        writer.writerow([r[c] for c in cols])
    fname = f"outreach_{deal['name'].replace(' ', '_') if deal else deal_id}.csv"
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


# ── Outreach ───────────────────────────────────────────────────────────────────

@app.route("/deals/<int:deal_id>/outreach/new", methods=["GET", "POST"])
def new_outreach(deal_id):
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger"); return redirect(url_for("deals"))
    if request.method == "POST":
        conn.execute(
            "INSERT INTO outreach (deal_id, contact_name, company, email, sent_by, "
            "outreach_date, outreach_summary, response_date, response_summary, follow_up_needed, follow_up_date, "
            "stage, interest_level) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (deal_id, request.form["contact_name"], request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("sent_by", ""),
             request.form.get("outreach_date") or None, request.form.get("outreach_summary", ""),
             request.form.get("response_date") or None, request.form.get("response_summary", ""),
             1 if request.form.get("follow_up_needed") else 0,
             request.form.get("follow_up_date") or None,
             request.form.get("stage", "Initial Email"), request.form.get("interest_level", "Unknown")),
        )
        conn.commit(); conn.close()
        flash("Outreach logged.", "success")
        return redirect(url_for("deal_detail", deal_id=deal_id))
    conn.close()
    return render_template("outreach_form.html", deal=deal, outreach=None, title="Log Outreach")


@app.route("/deals/<int:deal_id>/outreach/<int:outreach_id>/edit", methods=["GET", "POST"])
def edit_outreach(deal_id, outreach_id):
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    outreach = conn.execute("SELECT * FROM outreach WHERE id = ? AND deal_id = ?", (outreach_id, deal_id)).fetchone()
    if not deal or not outreach:
        conn.close(); flash("Not found.", "danger"); return redirect(url_for("deals"))
    if request.method == "POST":
        conn.execute(
            "UPDATE outreach SET contact_name=?, company=?, email=?, sent_by=?, "
            "outreach_date=?, outreach_summary=?, response_date=?, response_summary=?, "
            "follow_up_needed=?, follow_up_date=?, stage=?, interest_level=? WHERE id=?",
            (request.form["contact_name"], request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("sent_by", ""),
             request.form.get("outreach_date") or None, request.form.get("outreach_summary", ""),
             request.form.get("response_date") or None, request.form.get("response_summary", ""),
             1 if request.form.get("follow_up_needed") else 0,
             request.form.get("follow_up_date") or None,
             request.form.get("stage", "Initial Email"), request.form.get("interest_level", "Unknown"),
             outreach_id),
        )
        conn.commit(); conn.close()
        flash("Outreach updated.", "success")
        return redirect(url_for("deal_detail", deal_id=deal_id))
    conn.close()
    return render_template("outreach_form.html", deal=deal, outreach=outreach, title="Edit Outreach")


@app.route("/deals/<int:deal_id>/outreach/<int:outreach_id>/delete", methods=["POST"])
def delete_outreach(deal_id, outreach_id):
    conn = get_db()
    conn.execute("DELETE FROM outreach WHERE id = ? AND deal_id = ?", (outreach_id, deal_id))
    conn.commit(); conn.close()
    flash("Outreach entry deleted.", "info")
    return redirect(url_for("deal_detail", deal_id=deal_id))


@app.route("/deals/<int:deal_id>/outreach/<int:outreach_id>/log-response", methods=["POST"])
def log_outreach_response(deal_id, outreach_id):
    conn = get_db()
    conn.execute(
        "UPDATE outreach SET response_date=?, response_summary=?, interest_level=?, "
        "follow_up_needed=?, follow_up_date=?, stage=? WHERE id=? AND deal_id=?",
        (request.form.get("response_date") or None,
         request.form.get("response_summary", "").strip(),
         request.form.get("interest_level", "Unknown"),
         1 if request.form.get("follow_up_needed") else 0,
         request.form.get("follow_up_date") or None,
         request.form.get("stage", "Initial Email"),
         outreach_id, deal_id),
    )
    conn.commit(); conn.close()
    flash("Response logged.", "success")
    return redirect(url_for("deal_detail", deal_id=deal_id))


@app.route("/outreach/from-email", methods=["POST"])
def outreach_from_email():
    """One-click: log an outreach entry directly from an email."""
    deal_id = request.form.get("deal_id")
    if not deal_id:
        flash("Please select a deal first.", "warning")
        return redirect(request.referrer or url_for("emails"))
    conn = get_db()
    deal = conn.execute("SELECT id FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger")
        return redirect(request.referrer or url_for("emails"))
    conn.execute(
        "INSERT INTO outreach (deal_id, contact_name, company, email, sent_by, "
        "outreach_date, outreach_summary, follow_up_needed, stage, interest_level) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (deal_id,
         request.form.get("contact_name", ""),
         request.form.get("company", ""),
         request.form.get("contact_email", ""),
         request.form.get("sent_by", ""),
         request.form.get("outreach_date") or None,
         request.form.get("outreach_summary", ""),
         1 if request.form.get("follow_up_needed") else 0,
         request.form.get("stage", "Initial Email"),
         request.form.get("interest_level", "Unknown")),
    )
    conn.commit(); conn.close()
    flash("Outreach logged from email.", "success")
    return redirect(url_for("deal_detail", deal_id=deal_id))


@app.route("/deals/<int:deal_id>/report")
def deal_report(deal_id):
    """Printable one-pager for a deal."""
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger"); return redirect(url_for("deals"))
    outreaches = conn.execute(
        "SELECT * FROM outreach WHERE deal_id = ? ORDER BY outreach_date DESC",
        (deal_id,),
    ).fetchall()
    all_gps = conn.execute("SELECT * FROM gp_contacts").fetchall()
    matches = []
    for gp in all_gps:
        sc, reasons = score_gp(deal, gp)
        if sc > 0:
            matches.append({"gp": gp, "score": sc, "reasons": reasons})
    matches.sort(key=lambda x: x["score"], reverse=True)
    conn.close()
    from datetime import date as _date
    return render_template("deal_report.html", deal=deal, outreaches=outreaches,
                           matching_gps=matches[:10], now=_date.today().strftime("%B %d, %Y"))


# ── Activity Log ───────────────────────────────────────────────────────────────

@app.route("/activity")
def activity():
    conn = get_db()
    entries = conn.execute(
        "SELECT o.*, d.name as deal_name FROM outreach o "
        "JOIN deals d ON o.deal_id = d.id "
        "ORDER BY COALESCE(o.outreach_date, o.created_at) DESC LIMIT 150"
    ).fetchall()
    conn.close()
    return render_template("activity.html", entries=entries)


# ── GP Contacts ────────────────────────────────────────────────────────────────

@app.route("/contacts")
def contacts():
    q = request.args.get("q", "").strip()
    sf = request.args.get("status", "")
    conn = get_db()
    sql = "SELECT * FROM gp_contacts WHERE 1=1"
    params = []
    if q:
        sql += " AND (name LIKE ? OR company LIKE ? OR strategy LIKE ? OR location LIKE ?)"
        params += [f"%{q}%"] * 4
    if sf:
        sql += " AND status = ?"
        params.append(sf)
    rows = conn.execute(sql + " ORDER BY created_at DESC", params).fetchall()
    conn.close()
    return render_template("contacts.html", contacts=rows, q=q, status_filter=sf)


@app.route("/contacts/new", methods=["GET", "POST"])
def new_contact():
    if request.method == "POST":
        conn = get_db()
        conn.execute(
            "INSERT INTO gp_contacts (name, title, company, email, phone, location, "
            "aum_range, strategy, min_check, max_check, source, status, notes, last_contacted, next_contact_date) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""), request.form.get("location", ""),
             request.form.get("aum_range", ""), request.form.get("strategy", ""),
             request.form.get("min_check", ""), request.form.get("max_check", ""),
             request.form.get("source", ""), request.form.get("status", "New"),
             request.form.get("notes", ""), request.form.get("last_contacted") or None,
             request.form.get("next_contact_date") or None),
        )
        conn.commit(); conn.close()
        flash("GP contact added.", "success")
        return redirect(url_for("contacts"))
    return render_template("contact_form.html", contact=None, title="Add GP Contact")


@app.route("/contacts/<int:contact_id>")
def contact_detail(contact_id):
    conn = get_db()
    contact = conn.execute("SELECT * FROM gp_contacts WHERE id = ?", (contact_id,)).fetchone()
    if not contact:
        conn.close(); flash("Contact not found.", "danger"); return redirect(url_for("contacts"))
    notes = conn.execute(
        "SELECT * FROM contact_notes WHERE contact_id = ? ORDER BY note_date DESC, created_at DESC",
        (contact_id,),
    ).fetchall()
    conn.close()
    return render_template("contact_detail.html", contact=contact, notes=notes)


@app.route("/contacts/<int:contact_id>/edit", methods=["GET", "POST"])
def edit_contact(contact_id):
    conn = get_db()
    contact = conn.execute("SELECT * FROM gp_contacts WHERE id = ?", (contact_id,)).fetchone()
    if not contact:
        conn.close(); flash("Contact not found.", "danger"); return redirect(url_for("contacts"))
    if request.method == "POST":
        conn.execute(
            "UPDATE gp_contacts SET name=?, title=?, company=?, email=?, phone=?, location=?, "
            "aum_range=?, strategy=?, min_check=?, max_check=?, source=?, status=?, notes=?, "
            "last_contacted=?, next_contact_date=? WHERE id=?",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""), request.form.get("location", ""),
             request.form.get("aum_range", ""), request.form.get("strategy", ""),
             request.form.get("min_check", ""), request.form.get("max_check", ""),
             request.form.get("source", ""), request.form.get("status", "New"),
             request.form.get("notes", ""), request.form.get("last_contacted") or None,
             request.form.get("next_contact_date") or None, contact_id),
        )
        conn.commit(); conn.close()
        flash("Contact updated.", "success")
        return redirect(url_for("contact_detail", contact_id=contact_id))
    conn.close()
    return render_template("contact_form.html", contact=contact, title="Edit GP Contact")


@app.route("/contacts/<int:contact_id>/delete", methods=["POST"])
def delete_contact(contact_id):
    conn = get_db()
    conn.execute("DELETE FROM contact_notes WHERE contact_id = ?", (contact_id,))
    conn.execute("DELETE FROM gp_contacts WHERE id = ?", (contact_id,))
    conn.commit(); conn.close()
    flash("Contact deleted.", "info")
    return redirect(url_for("contacts"))


# ── Contact Quick Notes ────────────────────────────────────────────────────────

@app.route("/contacts/<int:contact_id>/notes/add", methods=["POST"])
def add_contact_note(contact_id):
    note_text = request.form.get("note_text", "").strip()
    if note_text:
        conn = get_db()
        conn.execute(
            "INSERT INTO contact_notes (contact_id, note_date, note_text) VALUES (?, ?, ?)",
            (contact_id, request.form.get("note_date") or date.today().isoformat(), note_text),
        )
        conn.commit(); conn.close()
        flash("Note added.", "success")
    return redirect(url_for("contact_detail", contact_id=contact_id))


@app.route("/contacts/<int:contact_id>/notes/<int:note_id>/delete", methods=["POST"])
def delete_contact_note(contact_id, note_id):
    conn = get_db()
    conn.execute("DELETE FROM contact_notes WHERE id = ? AND contact_id = ?", (note_id, contact_id))
    conn.commit(); conn.close()
    flash("Note deleted.", "info")
    return redirect(url_for("contact_detail", contact_id=contact_id))


@app.route("/company")
def company_view():
    name = request.args.get("name", "").strip()
    if not name:
        return redirect(url_for("contacts"))
    conn = get_db()
    gp_contacts = conn.execute(
        "SELECT * FROM gp_contacts WHERE company = ? ORDER BY name", (name,)
    ).fetchall()
    lp_investors = conn.execute(
        "SELECT * FROM lp_investors WHERE company = ? ORDER BY name", (name,)
    ).fetchall()

    # Gather all activity: contact_notes for GPs, lp_investor_notes for LPs
    activity = []
    for c in gp_contacts:
        rows = conn.execute(
            "SELECT *, 'gp' as kind, ? as person_name, ? as person_id FROM contact_notes WHERE contact_id = ?",
            (c["name"], c["id"], c["id"])
        ).fetchall()
        activity.extend(rows)
    for inv in lp_investors:
        rows = conn.execute(
            "SELECT *, 'lp' as kind, ? as person_name, ? as person_id FROM lp_investor_notes WHERE investor_id = ?",
            (inv["name"], inv["id"], inv["id"])
        ).fetchall()
        activity.extend(rows)

    # Sort activity by date desc
    activity.sort(key=lambda r: (r["note_date"] or ""), reverse=True)
    conn.close()
    return render_template("company.html", company_name=name,
                           gp_contacts=gp_contacts, lp_investors=lp_investors,
                           activity=activity)


# ── Email Templates ────────────────────────────────────────────────────────────

@app.route("/templates")
def templates():
    conn = get_db()
    rows = conn.execute("SELECT * FROM email_templates ORDER BY created_at").fetchall()
    conn.close()
    return render_template("templates.html", templates=rows)


@app.route("/templates/new", methods=["GET", "POST"])
def new_template():
    if request.method == "POST":
        conn = get_db()
        conn.execute(
            "INSERT INTO email_templates (name, subject, body) VALUES (?, ?, ?)",
            (request.form["name"], request.form.get("subject", ""), request.form["body"]),
        )
        conn.commit(); conn.close()
        flash("Template saved.", "success")
        return redirect(url_for("templates"))
    return render_template("template_form.html", template=None, title="New Template")


@app.route("/templates/<int:template_id>/edit", methods=["GET", "POST"])
def edit_template(template_id):
    conn = get_db()
    tmpl = conn.execute("SELECT * FROM email_templates WHERE id = ?", (template_id,)).fetchone()
    if not tmpl:
        conn.close(); flash("Not found.", "danger"); return redirect(url_for("templates"))
    if request.method == "POST":
        conn.execute(
            "UPDATE email_templates SET name=?, subject=?, body=? WHERE id=?",
            (request.form["name"], request.form.get("subject", ""), request.form["body"], template_id),
        )
        conn.commit(); conn.close()
        flash("Template updated.", "success")
        return redirect(url_for("templates"))
    conn.close()
    return render_template("template_form.html", template=tmpl, title="Edit Template")


@app.route("/templates/<int:template_id>/delete", methods=["POST"])
def delete_template(template_id):
    conn = get_db()
    conn.execute("DELETE FROM email_templates WHERE id = ?", (template_id,))
    conn.commit(); conn.close()
    flash("Template deleted.", "info")
    return redirect(url_for("templates"))


# ── Outlook Sync ───────────────────────────────────────────────────────────────

@app.route("/deals/<int:deal_id>/outlook", methods=["GET", "POST"])
def outlook_sync(deal_id):
    conn = get_db()
    deal = conn.execute("SELECT * FROM deals WHERE id = ?", (deal_id,)).fetchone()
    if not deal:
        conn.close(); flash("Deal not found.", "danger"); return redirect(url_for("deals"))

    emails = []
    error = None
    search_email = request.form.get("search_email", "").strip()
    days = int(request.form.get("days", 30))

    if request.method == "POST" and search_email:
        import threading

        result = {}

        def _fetch(search_email=search_email, days=days):
            try:
                import pythoncom
                import win32com.client
                pythoncom.CoInitialize()
                try:
                    outlook = win32com.client.Dispatch("Outlook.Application")
                    ns = outlook.GetNamespace("MAPI")
                    sent = ns.GetDefaultFolder(5)
                    cutoff = datetime.now() - timedelta(days=days)
                    found = []
                    for item in sent.Items:
                        try:
                            if item.Class != 43:
                                continue
                            sent_on = item.SentOn.replace(tzinfo=None)
                            if sent_on < cutoff:
                                continue
                            recipients = (item.To or "") + ";" + (item.CC or "")
                            if search_email.lower() not in recipients.lower():
                                continue
                            found.append({
                                "subject": item.Subject or "",
                                "to": item.To or "",
                                "sent_on": sent_on.strftime("%Y-%m-%d"),
                                "preview": (item.Body or "")[:400].replace("\r\n", " ").replace("\n", " "),
                            })
                        except Exception:
                            continue
                    result["emails"] = found
                finally:
                    pythoncom.CoUninitialize()
            except ImportError:
                result["error"] = "pywin32 not installed. In your Command Prompt run: pip install pywin32"
            except Exception as e:
                result["error"] = f"Could not connect to Outlook: {e}"

        t = threading.Thread(target=_fetch)
        t.start()
        t.join()

        if "error" in result:
            error = result["error"]
        else:
            emails = result.get("emails", [])

    conn.close()
    return render_template(
        "outlook_sync.html", deal=deal, emails=emails,
        error=error, search_email=search_email, days=days,
    )


@app.route("/deals/<int:deal_id>/outlook/import", methods=["POST"])
def outlook_import(deal_id):
    conn = get_db()
    conn.execute(
        "INSERT INTO outreach (deal_id, contact_name, company, email, sent_by, "
        "outreach_date, outreach_summary, follow_up_needed) VALUES (?, ?, ?, ?, ?, ?, ?, 0)",
        (deal_id, request.form.get("contact_name", ""), request.form.get("company", ""),
         request.form.get("email", ""), request.form.get("sent_by", ""),
         request.form.get("outreach_date", ""), request.form.get("outreach_summary", "")),
    )
    conn.commit(); conn.close()
    flash("Email imported as outreach entry.", "success")
    return redirect(url_for("deal_detail", deal_id=deal_id))


# ── Email Archive ──────────────────────────────────────────────────────────────

@app.route("/emails")
def emails():
    q = request.args.get("q", "").strip()
    deal_filter = request.args.get("deal_id", "")
    direction_filter = request.args.get("direction", "")
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    tag_filter = request.args.get("tag", "")

    conn = get_db()
    # Build base SQL with non-search filters
    sql = (
        "SELECT e.*, d.name as deal_name, g.name as gp_name "
        "FROM emails e "
        "LEFT JOIN deals d ON e.deal_id = d.id "
        "LEFT JOIN gp_contacts g ON e.contact_id = g.id "
        "WHERE 1=1"
    )
    params = []
    if deal_filter:
        sql += " AND e.deal_id = ?"
        params.append(deal_filter)
    if direction_filter:
        sql += " AND e.direction = ?"
        params.append(direction_filter)
    if date_from:
        sql += " AND e.sent_on >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND e.sent_on <= ?"
        params.append(date_to)
    if tag_filter:
        sql += " AND (',' || e.tags || ',') LIKE ?"
        params.append(f"%,{tag_filter},%")

    sql += " ORDER BY e.sent_on DESC LIMIT 500"
    rows = conn.execute(sql, params).fetchall()

    if q:
        try:
            from search_engine import semantic_search
            rows = [dict(r) for r in rows]
            results = semantic_search(
                q, rows,
                lambda e: " ".join(filter(None, [e.get("subject"), e.get("from_addr"), e.get("to_addr"), (e.get("body") or "")[:300]])),
                top_k=100, threshold=0.15,
            )
            rows = [item for item, _score in results]
        except Exception:
            # Fallback keyword search — handles "auto camp" matching "autocamp" etc.
            kw = q.lower().strip()
            kw_nospace = kw.replace(" ", "")
            words = [w for w in kw.split() if len(w) > 1]

            def _email_matches(r):
                fields = " ".join(filter(None, [
                    (r["subject"] or "").lower(),
                    (r["from_addr"] or "").lower(),
                    (r["to_addr"] or "").lower(),
                    (r["body"] or "")[:500].lower(),
                ]))
                fields_nospace = fields.replace(" ", "")
                return (kw in fields
                        or kw_nospace in fields_nospace
                        or (len(words) > 1 and all(w in fields for w in words)))

            rows = [r for r in rows if _email_matches(r)]
    selected_id = request.args.get("selected", type=int)
    deals = conn.execute("SELECT id, name FROM deals ORDER BY name").fetchall()
    contacts = conn.execute("SELECT id, name FROM gp_contacts ORDER BY name").fetchall()
    total = conn.execute("SELECT COUNT(*) FROM emails").fetchone()[0]
    raw_tags = conn.execute("SELECT tags FROM emails WHERE tags IS NOT NULL AND tags != ''").fetchall()
    all_tags = sorted(set(
        t.strip() for row in raw_tags for t in row["tags"].split(",") if t.strip()
    ))
    conn.close()

    # Determine which email to show in reading pane
    rows_list = list(rows)
    selected_email = None
    if selected_id:
        selected_email = next((r for r in rows_list if r["id"] == selected_id), None)
    if not selected_email and rows_list:
        selected_email = rows_list[0]

    return render_template(
        "emails.html", emails=rows_list, q=q, deals=deals, contacts=contacts,
        deal_filter=deal_filter, direction_filter=direction_filter,
        date_from=date_from, date_to=date_to, tag_filter=tag_filter,
        total=total, all_tags=all_tags, selected_email=selected_email,
    )


@app.route("/emails/<int:email_id>")
def email_detail(email_id):
    conn = get_db()
    email = conn.execute(
        "SELECT e.*, d.name as deal_name, g.name as gp_name "
        "FROM emails e "
        "LEFT JOIN deals d ON e.deal_id = d.id "
        "LEFT JOIN gp_contacts g ON e.contact_id = g.id "
        "WHERE e.id = ?", (email_id,)
    ).fetchone()
    deals = conn.execute("SELECT id, name FROM deals ORDER BY name").fetchall()
    contacts = conn.execute("SELECT id, name FROM gp_contacts ORDER BY name").fetchall()
    conn.close()
    if not email:
        flash("Email not found.", "danger"); return redirect(url_for("emails"))
    return render_template("email_detail.html", email=email, deals=deals, contacts=contacts)


@app.route("/emails/<int:email_id>/pane")
def email_pane(email_id):
    """Returns just the reading pane HTML fragment for the email archive AJAX loader."""
    conn = get_db()
    email = conn.execute(
        "SELECT e.*, d.name as deal_name, g.name as gp_name "
        "FROM emails e "
        "LEFT JOIN deals d ON e.deal_id = d.id "
        "LEFT JOIN gp_contacts g ON e.contact_id = g.id "
        "WHERE e.id = ?", (email_id,)
    ).fetchone()
    conn.close()
    if not email:
        return "<p class='text-danger p-3'>Email not found.</p>", 404
    return render_template("email_pane.html", e=email)


@app.route("/emails/sync", methods=["POST"])
def sync_emails():
    days = int(request.form.get("days", 1))
    include_inbox = request.form.get("include_inbox") == "1"

    # Pull DB data before entering the COM thread
    conn = get_db()
    deals = conn.execute("SELECT id, name, email_keywords FROM deals").fetchall()
    contacts = conn.execute(
        "SELECT id, email FROM gp_contacts WHERE email IS NOT NULL AND email != ''"
    ).fetchall()
    conn.close()

    contact_email_map = {c["email"].lower().strip(): c["id"] for c in contacts}
    deal_keywords = []
    for d in deals:
        if d["email_keywords"] and d["email_keywords"].strip():
            # Use manually set keywords (comma-separated)
            words = [w.strip().lower() for w in d["email_keywords"].split(",") if w.strip()]
        else:
            # Fall back to auto-generating from deal name
            words = [w for w in d["name"].lower().split() if len(w) > 3]
        if words:
            deal_keywords.append((words, d["id"]))

    import threading
    result = {}

    def _sync(days=days, include_inbox=include_inbox,
               contact_email_map=contact_email_map, deal_keywords=deal_keywords):
        try:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            try:
                outlook = win32com.client.Dispatch("Outlook.Application")
                ns = outlook.GetNamespace("MAPI")
                cutoff = datetime.now() - timedelta(days=days)

                own_addrs = set()
                try:
                    for acct in ns.Accounts:
                        own_addrs.add(acct.SmtpAddress.lower().strip())
                except Exception:
                    pass

                rows = []
                # Use start of the cutoff day so "Today only" includes all of today
                cutoff_day = cutoff.replace(hour=0, minute=0, second=0, microsecond=0)
                # %I = 12-hour clock (so midnight = "12:00 AM" not "00:00 AM" which Outlook rejects)
                cutoff_str = cutoff_day.strftime("%m/%d/%Y %I:%M %p")

                def _collect_folder(folder, direction):
                    try:
                        items = folder.Items
                        date_field = "[SentOn]" if direction == "sent" else "[ReceivedTime]"
                        # Sort newest-first before filtering so recent emails always come first
                        try:
                            items.Sort(date_field, True)
                        except Exception:
                            pass
                        try:
                            items = items.Restrict(f"{date_field} >= '{cutoff_str}'")
                        except Exception:
                            pass  # fallback: iterate all (slow but won't crash)
                        for item in items:
                            try:
                                if item.Class != 43:
                                    continue
                                sent_on = item.SentOn.replace(tzinfo=None) if hasattr(item, 'SentOn') else None
                                received_on = item.ReceivedTime.replace(tzinfo=None) if hasattr(item, 'ReceivedTime') else None
                                ts = sent_on or received_on
                                if not ts or ts < cutoff_day:
                                    continue
                                entry_id = item.EntryID
                                subject = item.Subject or ""
                                body = (item.Body or "")[:10000]
                                to_addr = item.To or ""
                                cc_addr = item.CC or ""
                                try:
                                    # Use MAPI property to get SMTP address directly —
                                    # avoids slow GetExchangeUser() Exchange round-trip
                                    PR_SMTP = "http://schemas.microsoft.com/mapi/proptag/0x39FE001E"
                                    from_addr = item.PropertyAccessor.GetProperty(PR_SMTP)
                                    if not from_addr:
                                        raise ValueError("empty")
                                except Exception:
                                    try:
                                        from_addr = item.SenderEmailAddress or item.SenderName or ""
                                    except Exception:
                                        from_addr = item.SenderName or ""
                                sent_on_str = ts.strftime("%Y-%m-%d %H:%M")

                                contact_id = None
                                all_addrs = (to_addr + ";" + cc_addr + ";" + from_addr).lower().replace(",", ";")
                                for addr in all_addrs.split(";"):
                                    addr = addr.strip()
                                    if addr and addr not in own_addrs and addr in contact_email_map:
                                        contact_id = contact_email_map[addr]
                                        break

                                deal_id = None
                                subject_lower = subject.lower()
                                for words, d_id in deal_keywords:
                                    if any(w in subject_lower for w in words):
                                        deal_id = d_id
                                        break

                                rows.append((subject, body, from_addr, to_addr, cc_addr,
                                             sent_on_str, direction, deal_id, contact_id, entry_id))
                            except Exception:
                                continue
                    except Exception:
                        pass

                _collect_folder(ns.GetDefaultFolder(5), "sent")
                if include_inbox:
                    _collect_folder(ns.GetDefaultFolder(6), "received")

                result["rows"] = rows
            finally:
                pythoncom.CoUninitialize()
        except ImportError:
            result["error"] = "pywin32 not installed. Run: pip install pywin32"
        except Exception as e:
            result["error"] = f"Could not connect to Outlook: {e}"

    t = threading.Thread(target=_sync)
    t.start()
    t.join(timeout=120)  # give up after 2 minutes
    if t.is_alive():
        result["error"] = "Sync timed out after 2 minutes. Make sure Outlook is open and try a shorter date range."

    if "error" in result:
        flash(result["error"], "danger")
    else:
        conn = get_db()
        synced = skipped = 0
        for row in result.get("rows", []):
            try:
                conn.execute(
                    "INSERT INTO emails (subject, body, from_addr, to_addr, cc_addr, sent_on, "
                    "direction, deal_id, contact_id, entry_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?) "
                    "ON CONFLICT(entry_id) DO NOTHING",
                    row,
                )
                if conn.execute("SELECT changes()").fetchone()[0]:
                    synced += 1
                else:
                    skipped += 1
            except Exception:
                skipped += 1
        conn.commit()
        conn.close()
        flash(f"Synced {synced} new emails ({skipped} already in archive).", "success")

    return redirect(url_for("emails"))


@app.route("/emails/<int:email_id>/link", methods=["POST"])
def link_email(email_id):
    deal_id = request.form.get("deal_id") or None
    contact_id = request.form.get("contact_id") or None
    tags = request.form.get("tags", "").strip()
    conn = get_db()
    conn.execute(
        "UPDATE emails SET deal_id=?, contact_id=?, tags=? WHERE id=?",
        (deal_id, contact_id, tags or None, email_id),
    )
    conn.commit()
    conn.close()
    flash("Email updated.", "success")
    return redirect(url_for("email_detail", email_id=email_id))


@app.route("/emails/<int:email_id>/delete", methods=["POST"])
def delete_email(email_id):
    conn = get_db()
    conn.execute("DELETE FROM emails WHERE id = ?", (email_id,))
    conn.commit(); conn.close()
    flash("Email removed from archive.", "info")
    return redirect(url_for("emails"))


# ── Reminders ──────────────────────────────────────────────────────────────────

@app.route("/reminders")
def reminders():
    today = date.today().isoformat()
    conn = get_db()
    outreach_reminders = conn.execute(
        "SELECT o.*, d.name as deal_name FROM outreach o "
        "JOIN deals d ON o.deal_id = d.id "
        "WHERE o.follow_up_needed = 1 "
        "ORDER BY COALESCE(o.follow_up_date, '9999') ASC",
    ).fetchall()
    contact_reminders = conn.execute(
        "SELECT * FROM gp_contacts WHERE next_contact_date IS NOT NULL AND next_contact_date != '' "
        "ORDER BY next_contact_date ASC",
    ).fetchall()
    all_contacts = conn.execute("SELECT id, name, company FROM gp_contacts ORDER BY name").fetchall()
    all_deals = conn.execute("SELECT id, name FROM deals ORDER BY name").fetchall()
    conn.close()
    return render_template(
        "reminders.html",
        outreach_reminders=outreach_reminders,
        contact_reminders=contact_reminders,
        all_contacts=all_contacts,
        all_deals=all_deals,
        today=today,
    )


@app.route("/reminders/quick", methods=["POST"])
def quick_reminder():
    """Quick follow-up reminder — can link to a GP contact or a deal outreach."""
    contact_id = request.form.get("contact_id") or None
    deal_id = request.form.get("deal_id") or None
    contact_name = request.form.get("contact_name", "").strip()
    reminder_date = request.form.get("reminder_date") or None
    note = request.form.get("note", "").strip()

    conn = get_db()

    if contact_id:
        # Set next_contact_date on the GP contact
        conn.execute(
            "UPDATE gp_contacts SET next_contact_date=? WHERE id=?",
            (reminder_date, contact_id),
        )
        if note:
            conn.execute(
                "INSERT INTO contact_notes (contact_id, note_date, note_text) VALUES (?, ?, ?)",
                (contact_id, date.today().isoformat(), f"[Reminder] {note}"),
            )
        conn.commit(); conn.close()
        flash(f"Reminder set for {reminder_date}.", "success")
    elif deal_id and contact_name:
        # Create an outreach entry with follow-up flag
        conn.execute(
            "INSERT INTO outreach (deal_id, contact_name, outreach_date, outreach_summary, "
            "follow_up_needed, follow_up_date, stage) VALUES (?, ?, ?, ?, 1, ?, 'Initial Email')",
            (deal_id, contact_name, date.today().isoformat(), note or "Follow-up reminder", reminder_date),
        )
        conn.commit(); conn.close()
        flash(f"Follow-up reminder set for {contact_name}.", "success")
    else:
        conn.close()
        flash("Please select a contact or deal.", "danger")

    next_url = request.form.get("next") or url_for("reminders")
    return redirect(next_url)


@app.route("/reminders/dismiss/<string:kind>/<int:rid>", methods=["POST"])
def dismiss_reminder(kind, rid):
    conn = get_db()
    if kind == "outreach":
        conn.execute("UPDATE outreach SET follow_up_needed=0, follow_up_date=NULL WHERE id=?", (rid,))
    elif kind == "contact":
        conn.execute("UPDATE gp_contacts SET next_contact_date=NULL WHERE id=?", (rid,))
    conn.commit(); conn.close()
    flash("Reminder dismissed.", "info")
    return redirect(request.referrer or url_for("reminders"))


# ── Smart Search ───────────────────────────────────────────────────────────────

@app.route("/search")
def search():
    q = request.args.get("q", "").strip()
    if not q:
        return render_template("search_results.html", q="", contacts=[], deals=[], emails=[], error=None)

    conn = get_db()
    contact_results = deal_results = email_results = []
    error = None

    try:
        from search_engine import semantic_search

        all_contacts = [dict(r) for r in conn.execute("SELECT * FROM gp_contacts").fetchall()]
        contact_results = semantic_search(
            q, all_contacts,
            lambda c: " ".join(filter(None, [
                c.get("name"), c.get("company"), c.get("title"),
                c.get("strategy"), c.get("location"), c.get("notes"),
                c.get("aum_range"), c.get("source"),
            ])),
            top_k=6,
        )

        all_deals = [dict(r) for r in conn.execute("SELECT * FROM deals").fetchall()]
        deal_results = semantic_search(
            q, all_deals,
            lambda d: " ".join(filter(None, [
                d.get("name"), d.get("asset_type"), d.get("location"),
                d.get("status"), d.get("notes"),
            ])),
            top_k=6,
        )

        all_emails = [dict(r) for r in conn.execute(
            "SELECT * FROM emails ORDER BY sent_on DESC LIMIT 500"
        ).fetchall()]
        email_results = semantic_search(
            q, all_emails,
            lambda e: " ".join(filter(None, [
                e.get("subject"), e.get("from_addr"), e.get("to_addr"),
            ])),
            top_k=6,
        )

    except ImportError:
        error = "Smart search not installed yet. Run: pip install sentence-transformers numpy"
    except Exception as ex:
        error = f"Search error: {ex}"

    conn.close()
    return render_template(
        "search_results.html",
        q=q,
        contacts=contact_results,
        deals=deal_results,
        emails=email_results,
        error=error,
    )


# ── LP Investors ───────────────────────────────────────────────────────────────

@app.route("/lp-investors")
def lp_investors():
    q = request.args.get("q", "").strip()
    sf = request.args.get("status", "")
    conn = get_db()
    sql = "SELECT * FROM lp_investors WHERE 1=1"
    params = []
    if q:
        sql += " AND (name LIKE ? OR company LIKE ? OR preferred_asset_types LIKE ? OR preferred_markets LIKE ?)"
        params += [f"%{q}%"] * 4
    if sf:
        sql += " AND status = ?"
        params.append(sf)
    rows = conn.execute(sql + " ORDER BY created_at DESC", params).fetchall()
    conn.close()
    return render_template("lp_investors.html", investors=rows, q=q, status_filter=sf)


@app.route("/lp-investors/new", methods=["GET", "POST"])
def new_lp_investor():
    if request.method == "POST":
        conn = get_db()
        conn.execute(
            "INSERT INTO lp_investors (name, title, company, email, phone, location, "
            "net_worth, preferred_asset_types, preferred_markets, min_check, max_check, "
            "hold_period_pref, accredited, source, status, notes, last_contacted, next_contact_date) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""), request.form.get("location", ""),
             request.form.get("net_worth", ""), request.form.get("preferred_asset_types", ""),
             request.form.get("preferred_markets", ""), request.form.get("min_check", ""),
             request.form.get("max_check", ""), request.form.get("hold_period_pref", ""),
             request.form.get("accredited", "Yes"), request.form.get("source", ""),
             request.form.get("status", "New"), request.form.get("notes", ""),
             request.form.get("last_contacted") or None, request.form.get("next_contact_date") or None),
        )
        conn.commit(); conn.close()
        flash("LP investor added.", "success")
        return redirect(url_for("lp_investors"))
    return render_template("lp_investor_form.html", investor=None, title="Add LP Investor")


@app.route("/lp-investors/<int:investor_id>")
def lp_investor_detail(investor_id):
    conn = get_db()
    investor = conn.execute("SELECT * FROM lp_investors WHERE id = ?", (investor_id,)).fetchone()
    if not investor:
        conn.close(); flash("Investor not found.", "danger"); return redirect(url_for("lp_investors"))
    notes = conn.execute(
        "SELECT * FROM lp_investor_notes WHERE investor_id = ? ORDER BY note_date DESC, created_at DESC",
        (investor_id,)
    ).fetchall()
    conn.close()
    return render_template("lp_investor_detail.html", investor=investor, notes=notes)


@app.route("/lp-investors/<int:investor_id>/notes/add", methods=["POST"])
def add_lp_investor_note(investor_id):
    conn = get_db()
    note_date = request.form.get("note_date") or date.today().isoformat()
    note_text = request.form.get("note_text", "").strip()
    if note_text:
        conn.execute(
            "INSERT INTO lp_investor_notes (investor_id, note_date, note_text) VALUES (?, ?, ?)",
            (investor_id, note_date, note_text),
        )
        conn.commit()
    conn.close()
    return redirect(url_for("lp_investor_detail", investor_id=investor_id))


@app.route("/lp-investors/<int:investor_id>/notes/<int:note_id>/delete", methods=["POST"])
def delete_lp_investor_note(investor_id, note_id):
    conn = get_db()
    conn.execute("DELETE FROM lp_investor_notes WHERE id = ? AND investor_id = ?", (note_id, investor_id))
    conn.commit()
    conn.close()
    return redirect(url_for("lp_investor_detail", investor_id=investor_id))


@app.route("/lp-investors/<int:investor_id>/edit", methods=["GET", "POST"])
def edit_lp_investor(investor_id):
    conn = get_db()
    investor = conn.execute("SELECT * FROM lp_investors WHERE id = ?", (investor_id,)).fetchone()
    if not investor:
        conn.close(); flash("Investor not found.", "danger"); return redirect(url_for("lp_investors"))
    if request.method == "POST":
        conn.execute(
            "UPDATE lp_investors SET name=?, title=?, company=?, email=?, phone=?, location=?, "
            "net_worth=?, preferred_asset_types=?, preferred_markets=?, min_check=?, max_check=?, "
            "hold_period_pref=?, accredited=?, source=?, status=?, notes=?, last_contacted=?, "
            "next_contact_date=? WHERE id=?",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""), request.form.get("location", ""),
             request.form.get("net_worth", ""), request.form.get("preferred_asset_types", ""),
             request.form.get("preferred_markets", ""), request.form.get("min_check", ""),
             request.form.get("max_check", ""), request.form.get("hold_period_pref", ""),
             request.form.get("accredited", "Yes"), request.form.get("source", ""),
             request.form.get("status", "New"), request.form.get("notes", ""),
             request.form.get("last_contacted") or None, request.form.get("next_contact_date") or None,
             investor_id),
        )
        conn.commit(); conn.close()
        flash("LP investor updated.", "success")
        return redirect(url_for("lp_investor_detail", investor_id=investor_id))
    conn.close()
    return render_template("lp_investor_form.html", investor=investor, title="Edit LP Investor")


@app.route("/lp-investors/<int:investor_id>/delete", methods=["POST"])
def delete_lp_investor(investor_id):
    conn = get_db()
    conn.execute("DELETE FROM lp_investors WHERE id = ?", (investor_id,))
    conn.commit(); conn.close()
    flash("LP investor deleted.", "info")
    return redirect(url_for("lp_investors"))


@app.route("/lp-investors/export.csv")
def export_lp_investors():
    conn = get_db()
    rows = conn.execute("SELECT * FROM lp_investors ORDER BY name").fetchall()
    conn.close()
    cols = ["id", "name", "title", "company", "email", "phone", "location", "net_worth",
            "preferred_asset_types", "preferred_markets", "min_check", "max_check",
            "hold_period_pref", "accredited", "source", "status", "notes",
            "last_contacted", "next_contact_date", "created_at"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cols)
    for r in rows:
        writer.writerow([r[c] for c in cols])
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=lp_investors.csv"})


# ── Lenders ────────────────────────────────────────────────────────────────────

@app.route("/lenders")
def lenders():
    q = request.args.get("q", "").strip()
    sf = request.args.get("status", "")
    lt = request.args.get("lender_type", "")
    conn = get_db()
    sql = "SELECT * FROM lenders WHERE 1=1"
    params = []
    if q:
        sql += " AND (name LIKE ? OR company LIKE ? OR asset_types LIKE ? OR markets LIKE ? OR loan_types LIKE ?)"
        params += [f"%{q}%"] * 5
    if sf:
        sql += " AND status = ?"
        params.append(sf)
    if lt:
        sql += " AND lender_type = ?"
        params.append(lt)
    rows = conn.execute(sql + " ORDER BY created_at DESC", params).fetchall()
    conn.close()
    return render_template("lenders.html", lenders=rows, q=q, status_filter=sf, lender_type_filter=lt)


@app.route("/lenders/new", methods=["GET", "POST"])
def new_lender():
    if request.method == "POST":
        conn = get_db()
        conn.execute(
            "INSERT INTO lenders (name, title, company, email, phone, lender_type, loan_types, "
            "asset_types, min_loan, max_loan, max_ltv, typical_rate, recourse, markets, "
            "typical_term, origination_fee, status, notes, last_contacted) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""),
             request.form.get("lender_type", ""), request.form.get("loan_types", ""),
             request.form.get("asset_types", ""), request.form.get("min_loan", ""),
             request.form.get("max_loan", ""), request.form.get("max_ltv", ""),
             request.form.get("typical_rate", ""), request.form.get("recourse", ""),
             request.form.get("markets", ""), request.form.get("typical_term", ""),
             request.form.get("origination_fee", ""), request.form.get("status", "New"),
             request.form.get("notes", ""), request.form.get("last_contacted") or None),
        )
        conn.commit(); conn.close()
        flash("Lender added.", "success")
        return redirect(url_for("lenders"))
    return render_template("lender_form.html", lender=None, title="Add Lender")


@app.route("/lenders/<int:lender_id>")
def lender_detail(lender_id):
    conn = get_db()
    lender = conn.execute("SELECT * FROM lenders WHERE id = ?", (lender_id,)).fetchone()
    conn.close()
    if not lender:
        flash("Lender not found.", "danger"); return redirect(url_for("lenders"))
    return render_template("lender_detail.html", lender=lender)


@app.route("/lenders/<int:lender_id>/edit", methods=["GET", "POST"])
def edit_lender(lender_id):
    conn = get_db()
    lender = conn.execute("SELECT * FROM lenders WHERE id = ?", (lender_id,)).fetchone()
    if not lender:
        conn.close(); flash("Lender not found.", "danger"); return redirect(url_for("lenders"))
    if request.method == "POST":
        conn.execute(
            "UPDATE lenders SET name=?, title=?, company=?, email=?, phone=?, lender_type=?, "
            "loan_types=?, asset_types=?, min_loan=?, max_loan=?, max_ltv=?, typical_rate=?, "
            "recourse=?, markets=?, typical_term=?, origination_fee=?, status=?, notes=?, "
            "last_contacted=? WHERE id=?",
            (request.form["name"], request.form.get("title", ""), request.form.get("company", ""),
             request.form.get("email", ""), request.form.get("phone", ""),
             request.form.get("lender_type", ""), request.form.get("loan_types", ""),
             request.form.get("asset_types", ""), request.form.get("min_loan", ""),
             request.form.get("max_loan", ""), request.form.get("max_ltv", ""),
             request.form.get("typical_rate", ""), request.form.get("recourse", ""),
             request.form.get("markets", ""), request.form.get("typical_term", ""),
             request.form.get("origination_fee", ""), request.form.get("status", "New"),
             request.form.get("notes", ""), request.form.get("last_contacted") or None, lender_id),
        )
        conn.commit(); conn.close()
        flash("Lender updated.", "success")
        return redirect(url_for("lender_detail", lender_id=lender_id))
    conn.close()
    return render_template("lender_form.html", lender=lender, title="Edit Lender")


@app.route("/lenders/<int:lender_id>/delete", methods=["POST"])
def delete_lender(lender_id):
    conn = get_db()
    conn.execute("DELETE FROM lenders WHERE id = ?", (lender_id,))
    conn.commit(); conn.close()
    flash("Lender deleted.", "info")
    return redirect(url_for("lenders"))


@app.route("/lenders/export.csv")
def export_lenders():
    conn = get_db()
    rows = conn.execute("SELECT * FROM lenders ORDER BY name").fetchall()
    conn.close()
    cols = ["id", "name", "title", "company", "email", "phone", "lender_type", "loan_types",
            "asset_types", "min_loan", "max_loan", "max_ltv", "typical_rate", "recourse",
            "markets", "typical_term", "origination_fee", "status", "notes",
            "last_contacted", "created_at"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cols)
    for r in rows:
        writer.writerow([r[c] for c in cols])
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=lenders.csv"})


# ── Email Campaigns ────────────────────────────────────────────────────────────

def _apply_merge(text, contact):
    """Replace {merge_fields} in text with contact data."""
    if not text:
        return text
    name = contact["name"] or ""
    first_name = name.split()[0] if name else ""
    company = contact["company"] or "your firm"
    strategy = contact["strategy"] or ""
    location = contact["location"] or ""
    aum = contact["aum_range"] or ""
    return (text
            .replace("{first_name}", first_name)
            .replace("{full_name}", name)
            .replace("{company}", company)
            .replace("{strategy}", strategy)
            .replace("{location}", location)
            .replace("{aum}", aum))


def _apply_merge_lp(text, investor):
    """Replace {merge_fields} in text with LP investor data."""
    if not text:
        return text
    name = investor["name"] or ""
    first_name = name.split()[0] if name else ""
    company = investor["company"] or "your firm"
    asset_types = investor["preferred_asset_types"] or ""
    markets = investor["preferred_markets"] or ""
    return (text
            .replace("{first_name}", first_name)
            .replace("{full_name}", name)
            .replace("{company}", company)
            .replace("{asset_types}", asset_types)
            .replace("{markets}", markets))


import re

def _email_ready_html(html):
    """Make Quill HTML email-friendly by removing paragraph spacing."""
    if not html:
        return html
    # Add margin:0 to <p> tags so email clients don't add extra spacing
    html = re.sub(r'<p(?=[\s>])', '<p style="margin:0;padding:0;"', html)
    return html


@app.route("/campaigns")
def campaigns():
    conn = get_db()
    rows = conn.execute(
        "SELECT c.*, "
        "(SELECT COUNT(*) FROM campaign_emails WHERE campaign_id=c.id) as total_count, "
        "(SELECT COUNT(*) FROM campaign_emails WHERE campaign_id=c.id AND status='sent') as sent_count "
        "FROM campaigns c ORDER BY c.created_at DESC"
    ).fetchall()
    conn.close()
    return render_template("campaigns.html", campaigns=rows)


@app.route("/campaigns/new", methods=["GET", "POST"])
def new_campaign():
    conn = get_db()
    templates_list = conn.execute("SELECT * FROM email_templates ORDER BY name").fetchall()
    gp_batches = conn.execute(
        "SELECT * FROM import_batches WHERE type='gp' ORDER BY created_at DESC"
    ).fetchall()
    lp_batches = conn.execute(
        "SELECT * FROM import_batches WHERE type='lp' ORDER BY created_at DESC"
    ).fetchall()
    if request.method == "POST":
        name = request.form["name"].strip()
        subject = request.form["subject"].strip()
        body = request.form["body"].strip()
        target_audience = request.form.get("target_audience", "gp")

        status_filter = request.form.getlist("status_filter")
        selected_batch_ids = request.form.getlist("batch_ids")
        lp_handpick_ids = request.form.getlist("lp_handpick_ids")
        gp_handpick_ids = request.form.getlist("gp_handpick_ids")

        if target_audience == "lp":
            if lp_handpick_ids:
                # Hand-picked individuals override filters
                placeholders = ",".join("?" * len(lp_handpick_ids))
                recipients = conn.execute(
                    f"SELECT * FROM lp_investors WHERE id IN ({placeholders}) ORDER BY name",
                    lp_handpick_ids,
                ).fetchall()
            else:
                # Filter LP investors
                sql = "SELECT * FROM lp_investors WHERE email IS NOT NULL AND email != '' AND email LIKE '%@%'"
                params = []
                if status_filter:
                    placeholders = ",".join("?" * len(status_filter))
                    sql += f" AND status IN ({placeholders})"
                    params += status_filter
                asset_type_kw = request.form.get("asset_type_kw", "").strip()
                if asset_type_kw:
                    sql += " AND preferred_asset_types LIKE ?"
                    params.append(f"%{asset_type_kw}%")
                if selected_batch_ids:
                    placeholders = ",".join("?" * len(selected_batch_ids))
                    sql += f" AND import_batch_id IN ({placeholders})"
                    params += selected_batch_ids
                sql += " ORDER BY name"
                recipients = conn.execute(sql, params).fetchall()

            if not recipients:
                conn.close()
                flash("No LP investors match those filters (or none have email addresses). Adjust filters and try again.", "warning")
                return render_template("campaign_new.html", templates=templates_list,
                                       gp_batches=gp_batches, lp_batches=lp_batches, title="New Campaign")

            cur = conn.execute(
                "INSERT INTO campaigns (name, subject, body_template, target_audience) VALUES (?, ?, ?, ?)",
                (name, subject, body, "lp"),
            )
            campaign_id = cur.lastrowid

            for inv in recipients:
                rendered_subject = _apply_merge_lp(subject, inv)
                rendered_body = _email_ready_html(_apply_merge_lp(body, inv))
                conn.execute(
                    "INSERT INTO campaign_emails (campaign_id, lp_investor_id, to_name, to_email, company, subject, body) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?)",
                    (campaign_id, inv["id"], inv["name"], inv["email"], inv["company"],
                     rendered_subject, rendered_body),
                )

        else:
            if gp_handpick_ids:
                # Hand-picked individuals override filters
                placeholders = ",".join("?" * len(gp_handpick_ids))
                recipients = conn.execute(
                    f"SELECT * FROM gp_contacts WHERE id IN ({placeholders}) ORDER BY name",
                    gp_handpick_ids,
                ).fetchall()
            else:
                # Filter GPs
                sql = "SELECT * FROM gp_contacts WHERE email IS NOT NULL AND email != '' AND email LIKE '%@%'"
                params = []
                if status_filter:
                    placeholders = ",".join("?" * len(status_filter))
                    sql += f" AND status IN ({placeholders})"
                    params += status_filter
                strategy_kw = request.form.get("strategy_kw", "").strip()
                if strategy_kw:
                    sql += " AND strategy LIKE ?"
                    params.append(f"%{strategy_kw}%")
                if selected_batch_ids:
                    placeholders = ",".join("?" * len(selected_batch_ids))
                    sql += f" AND import_batch_id IN ({placeholders})"
                    params += selected_batch_ids
                sql += " ORDER BY name"
                recipients = conn.execute(sql, params).fetchall()

            if not recipients:
                conn.close()
                flash("No GP contacts match those filters (or none have email addresses). Adjust filters and try again.", "warning")
                return render_template("campaign_new.html", templates=templates_list,
                                       gp_batches=gp_batches, lp_batches=lp_batches, title="New Campaign")

            cur = conn.execute(
                "INSERT INTO campaigns (name, subject, body_template, target_audience) VALUES (?, ?, ?, ?)",
                (name, subject, body, "gp"),
            )
            campaign_id = cur.lastrowid

            for gp in recipients:
                rendered_subject = _apply_merge(subject, gp)
                rendered_body = _email_ready_html(_apply_merge(body, gp))
                conn.execute(
                    "INSERT INTO campaign_emails (campaign_id, contact_id, to_name, to_email, company, subject, body) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?)",
                    (campaign_id, gp["id"], gp["name"], gp["email"], gp["company"],
                     rendered_subject, rendered_body),
                )

        conn.commit()
        conn.close()
        flash(f"Generated {len(recipients)} draft emails. Review them below then send.", "success")
        return redirect(url_for("campaign_review", campaign_id=campaign_id))

    conn.close()
    return render_template("campaign_new.html", templates=templates_list,
                           gp_batches=gp_batches, lp_batches=lp_batches, title="New Campaign")


@app.route("/campaigns/contacts-search")
def campaigns_contacts_search():
    """AJAX: search GP contacts or LP investors for hand-pick selection."""
    q = request.args.get("q", "").strip()
    audience = request.args.get("audience", "gp")
    like = f"%{q}%"
    conn = get_db()
    if audience == "lp":
        rows = conn.execute(
            "SELECT id, name, company, email FROM lp_investors "
            "WHERE (name LIKE ? OR company LIKE ? OR email LIKE ?) "
            "ORDER BY name LIMIT 60",
            (like, like, like),
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT id, name, company, email FROM gp_contacts "
            "WHERE (name LIKE ? OR company LIKE ? OR email LIKE ?) "
            "ORDER BY name LIMIT 60",
            (like, like, like),
        ).fetchall()
    conn.close()
    return {"results": [dict(r) for r in rows]}


@app.route("/campaigns/<int:campaign_id>/review")
def campaign_review(campaign_id):
    conn = get_db()
    campaign = conn.execute("SELECT * FROM campaigns WHERE id = ?", (campaign_id,)).fetchone()
    if not campaign:
        conn.close(); flash("Campaign not found.", "danger"); return redirect(url_for("campaigns"))
    emails_list = conn.execute(
        "SELECT ce.*, g.strategy, g.aum_range "
        "FROM campaign_emails ce "
        "LEFT JOIN gp_contacts g ON ce.contact_id = g.id "
        "WHERE ce.campaign_id = ? ORDER BY ce.to_name",
        (campaign_id,),
    ).fetchall()
    conn.close()
    pending = [e for e in emails_list if e["status"] == "pending"]
    sent = [e for e in emails_list if e["status"] == "sent"]
    return render_template("campaign_review.html", campaign=campaign,
                           emails=emails_list, pending=pending, sent=sent)


@app.route("/campaigns/<int:campaign_id>/send", methods=["POST"])
def campaign_send(campaign_id):
    conn = get_db()
    campaign = conn.execute("SELECT * FROM campaigns WHERE id = ?", (campaign_id,)).fetchone()
    if not campaign:
        conn.close(); flash("Campaign not found.", "danger"); return redirect(url_for("campaigns"))

    selected_ids = request.form.getlist("email_ids")
    if not selected_ids:
        conn.close(); flash("No emails selected.", "warning")
        return redirect(url_for("campaign_review", campaign_id=campaign_id))

    selected_ids = [int(i) for i in selected_ids]
    rows = conn.execute(
        f"SELECT * FROM campaign_emails WHERE campaign_id = ? AND id IN ({','.join('?'*len(selected_ids))})",
        [campaign_id] + selected_ids,
    ).fetchall()

    import threading
    result = {"sent": 0, "failed": 0, "error": None}

    def _send_all(rows=rows):
        try:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            try:
                outlook = win32com.client.Dispatch("Outlook.Application")
                namespace = outlook.GetNamespace("MAPI")
                sent_folder = namespace.GetDefaultFolder(5)  # 5 = olFolderSentMail
                for row in rows:
                    try:
                        mail = outlook.CreateItem(0)
                        mail.To = row["to_email"]
                        mail.Subject = row["subject"]
                        mail.HTMLBody = row["body"]
                        mail.DeleteAfterSubmit = False
                        mail.SaveSentMessageFolder = sent_folder
                        mail.Send()
                        result["sent"] += 1
                        result[f"ok_{row['id']}"] = True
                    except Exception as e:
                        result["failed"] += 1
                        result[f"err_{row['id']}"] = str(e)
            finally:
                pythoncom.CoUninitialize()
        except ImportError:
            result["error"] = "pywin32 not installed. Run: pip install pywin32"
        except Exception as e:
            result["error"] = f"Could not connect to Outlook: {e}"

    t = threading.Thread(target=_send_all)
    t.start()
    t.join(timeout=300)

    if result["error"]:
        conn.close()
        flash(result["error"], "danger")
        return redirect(url_for("campaign_review", campaign_id=campaign_id))

    today = date.today().isoformat()
    for row in rows:
        if result.get(f"ok_{row['id']}"):
            conn.execute(
                "UPDATE campaign_emails SET status='sent', sent_at=? WHERE id=?",
                (datetime.now().strftime("%Y-%m-%d %H:%M"), row["id"]),
            )
            # Update GP last_contacted
            if row["contact_id"]:
                conn.execute(
                    "UPDATE gp_contacts SET last_contacted=? WHERE id=?",
                    (today, row["contact_id"]),
                )
                conn.execute(
                    "INSERT INTO contact_notes (contact_id, note_date, note_text, note_html) VALUES (?, ?, ?, ?)",
                    (row["contact_id"], today,
                     f"[Campaign: {campaign['name']}] Email sent — Subject: {row['subject']}",
                     row["body"]),
                )
            # Update LP last_contacted
            if row["lp_investor_id"]:
                conn.execute(
                    "UPDATE lp_investors SET last_contacted=? WHERE id=?",
                    (today, row["lp_investor_id"]),
                )
                conn.execute(
                    "INSERT INTO lp_investor_notes (investor_id, note_date, note_text, note_html) VALUES (?, ?, ?, ?)",
                    (row["lp_investor_id"], today,
                     f"[Campaign: {campaign['name']}] Email sent — Subject: {row['subject']}",
                     row["body"]),
                )

    # Mark campaign as sent if all emails done
    pending_left = conn.execute(
        "SELECT COUNT(*) FROM campaign_emails WHERE campaign_id=? AND status='pending'",
        (campaign_id,),
    ).fetchone()[0]
    if pending_left == 0:
        conn.execute(
            "UPDATE campaigns SET status='sent', sent_at=? WHERE id=?",
            (datetime.now().strftime("%Y-%m-%d %H:%M"), campaign_id),
        )

    conn.commit()
    conn.close()

    if result["failed"]:
        flash(f"Sent {result['sent']} emails. {result['failed']} failed — check the review page.", "warning")
    else:
        flash(f"Sent {result['sent']} emails successfully.", "success")

    return redirect(url_for("campaign_review", campaign_id=campaign_id))


@app.route("/campaigns/<int:campaign_id>/delete", methods=["POST"])
def delete_campaign(campaign_id):
    conn = get_db()
    conn.execute("DELETE FROM campaign_emails WHERE campaign_id = ?", (campaign_id,))
    conn.execute("DELETE FROM campaigns WHERE id = ?", (campaign_id,))
    conn.commit(); conn.close()
    flash("Campaign deleted.", "info")
    return redirect(url_for("campaigns"))


# ── Preqin / Excel Import ───────────────────────────────────────────────────────

# Known Preqin column name variations → our field names
_PREQIN_COL_MAP = {
    "name":         ["contact name", "name", "full name", "contact", "first name"],
    "company":      ["fund manager", "manager name", "firm", "company", "company name",
                     "organization", "employer", "fund manager name", "manager", "firm name"],
    "title":        ["job title", "title", "position", "role"],
    "email":        ["email address", "email", "e-mail", "primary email"],
    "phone":        ["phone", "phone number", "telephone", "mobile"],
    "location":     ["city", "location", "country", "hq location", "office location", "geography"],
    "aum_range":    ["aum", "total aum", "aum (usd mn)", "assets under management", "aum range",
                     "firm aum", "fund size"],
    "strategy":     ["strategy", "investment strategy", "fund strategy", "asset class", "focus",
                     "primary strategy"],
    "min_check":    ["min check", "minimum check", "min equity", "minimum investment"],
    "max_check":    ["max check", "maximum check", "max equity", "maximum investment"],
    "notes":        ["notes", "comments", "description"],
}


def _detect_col(headers, candidates):
    """Return the first header that matches any candidate (case-insensitive)."""
    lower_headers = [h.lower().strip() for h in headers]
    for c in candidates:
        if c in lower_headers:
            return headers[lower_headers.index(c)]
    return None


@app.route("/contacts/import", methods=["GET", "POST"])
def import_contacts():
    if request.method == "GET":
        return render_template("contacts_import.html")

    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        flash("Please select a file to upload.", "warning")
        return render_template("contacts_import.html")

    fname = uploaded.filename.lower()
    rows = []
    headers = []

    try:
        if fname.endswith(".csv"):
            import io as _io
            content = uploaded.read().decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(_io.StringIO(content))
            headers = reader.fieldnames or []
            rows = list(reader)
        elif fname.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                flash("The uploaded file appears to be empty.", "warning")
                return render_template("contacts_import.html")
            headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
            for r in all_rows[1:]:
                rows.append(dict(zip(headers, [str(v).strip() if v is not None else "" for v in r])))
        else:
            flash("Please upload a .csv or .xlsx file.", "warning")
            return render_template("contacts_import.html")
    except Exception as e:
        flash(f"Could not read file: {e}", "danger")
        return render_template("contacts_import.html")

    # Build column mapping
    col_map = {field: _detect_col(headers, candidates) for field, candidates in _PREQIN_COL_MAP.items()}

    # Parse rows into contacts
    preview = []
    for r in rows:
        def g(field):
            col = col_map.get(field)
            return (r.get(col) or "").strip() if col else ""

        entry = {
            "name": g("name"), "company": g("company"), "title": g("title"),
            "email": g("email"), "phone": g("phone"), "location": g("location"),
            "aum_range": g("aum_range"), "strategy": g("strategy"),
            "min_check": g("min_check"), "max_check": g("max_check"), "notes": g("notes"),
        }
        if entry["name"] or entry["company"] or entry["email"]:
            preview.append(entry)

    if not preview:
        flash("No usable rows found. Make sure the file has contact name, company, or email columns.", "warning")
        return render_template("contacts_import.html")

    import_label = request.form.get("import_label", "").strip()
    return render_template("contacts_import.html", preview=preview, col_map=col_map,
                           row_count=len(preview), import_label=import_label)


@app.route("/contacts/import/confirm", methods=["POST"])
def import_contacts_confirm():
    """Receives the confirmed hidden-field rows and does the actual DB insert."""
    i = 0
    entries = []
    while request.form.get(f"rows[{i}][name]") is not None or request.form.get(f"rows[{i}][email]") is not None:
        entries.append({
            "name":      request.form.get(f"rows[{i}][name]", "").strip(),
            "company":   request.form.get(f"rows[{i}][company]", "").strip(),
            "title":     request.form.get(f"rows[{i}][title]", "").strip(),
            "email":     request.form.get(f"rows[{i}][email]", "").strip(),
            "phone":     request.form.get(f"rows[{i}][phone]", "").strip(),
            "location":  request.form.get(f"rows[{i}][location]", "").strip(),
            "aum_range": request.form.get(f"rows[{i}][aum_range]", "").strip(),
            "strategy":  request.form.get(f"rows[{i}][strategy]", "").strip(),
            "min_check": request.form.get(f"rows[{i}][min_check]", "").strip(),
            "max_check": request.form.get(f"rows[{i}][max_check]", "").strip(),
            "notes":     request.form.get(f"rows[{i}][notes]", "").strip(),
        })
        i += 1

    if not entries:
        flash("No data received. Please re-upload your file.", "warning")
        return redirect(url_for("import_contacts"))

    import_label = request.form.get("import_label", "").strip() or "GP Import"
    valid_entries = [e for e in entries if e["name"] or e["email"]]

    conn = get_db()
    cur = conn.execute(
        "INSERT INTO import_batches (label, type, count) VALUES (?, 'gp', ?)",
        (import_label, len(valid_entries)),
    )
    batch_id = cur.lastrowid

    imported = 0
    for entry in valid_entries:
        name = entry["name"] or entry["company"] or "Unknown"
        conn.execute(
            "INSERT INTO gp_contacts (name, title, company, email, phone, location, "
            "aum_range, strategy, min_check, max_check, source, status, notes, import_batch_id) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Preqin', 'New', ?, ?)",
            (name, entry["title"], entry["company"], entry["email"], entry["phone"],
             entry["location"], entry["aum_range"], entry["strategy"],
             entry["min_check"], entry["max_check"], entry["notes"], batch_id),
        )
        imported += 1
    conn.commit(); conn.close()
    flash(f"Imported {imported} GP contacts as '{import_label}'.", "success")
    return redirect(url_for("contacts"))


# ── LP Investor Excel / CSV Import ────────────────────────────────────────────

_LP_COL_MAP = {
    "name":                  ["contact name", "name", "full name", "investor name", "first name"],
    "company":               ["firm", "company", "company name", "organization", "fund manager",
                              "employer", "firm name", "investor firm", "lp firm",
                              "investor", "institution", "institution name", "plan sponsor",
                              "fund name", "lp name", "lp", "entity", "entity name"],
    "title":                 ["title", "job title", "position", "role"],
    "email":                 ["email", "email address", "e-mail", "primary email"],
    "phone":                 ["phone", "phone number", "telephone", "mobile"],
    "location":              ["city", "location", "country", "state", "hq location", "geography"],
    "net_worth":             ["net worth", "aum", "total aum", "assets under management",
                              "estimated net worth", "wealth"],
    "preferred_asset_types": ["asset type", "asset class", "preferred asset types",
                              "investment type", "focus", "strategy", "investment strategy"],
    "preferred_markets":     ["markets", "preferred markets", "geography", "target markets",
                              "market focus", "regions"],
    "min_check":             ["min check", "minimum check", "min investment",
                              "minimum investment", "min equity"],
    "max_check":             ["max check", "maximum check", "max investment",
                              "maximum investment", "max equity"],
    "notes":                 ["notes", "comments", "description"],
}


@app.route("/lp-investors/import", methods=["GET", "POST"])
def import_lp_investors():
    if request.method == "GET":
        return render_template("lp_investors_import.html")

    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        flash("Please select a file to upload.", "warning")
        return render_template("lp_investors_import.html")

    fname = uploaded.filename.lower()
    rows = []
    headers = []

    try:
        if fname.endswith(".csv"):
            import io as _io
            content = uploaded.read().decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(_io.StringIO(content))
            headers = reader.fieldnames or []
            rows = list(reader)
        elif fname.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                flash("The uploaded file appears to be empty.", "warning")
                return render_template("lp_investors_import.html")
            headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
            for r in all_rows[1:]:
                rows.append(dict(zip(headers, [str(v).strip() if v is not None else "" for v in r])))
        else:
            flash("Please upload a .csv or .xlsx file.", "warning")
            return render_template("lp_investors_import.html")
    except Exception as e:
        flash(f"Could not read file: {e}", "danger")
        return render_template("lp_investors_import.html")

    col_map = {field: _detect_col(headers, candidates) for field, candidates in _LP_COL_MAP.items()}

    preview = []
    for r in rows:
        def g(field, _r=r):
            col = col_map.get(field)
            return (_r.get(col) or "").strip() if col else ""

        entry = {
            "name": g("name"), "company": g("company"), "title": g("title"),
            "email": g("email"), "phone": g("phone"), "location": g("location"),
            "net_worth": g("net_worth"), "preferred_asset_types": g("preferred_asset_types"),
            "preferred_markets": g("preferred_markets"), "min_check": g("min_check"),
            "max_check": g("max_check"), "notes": g("notes"),
        }
        if entry["name"] or entry["company"] or entry["email"]:
            preview.append(entry)

    if not preview:
        flash("No usable rows found. Make sure the file has a name, company, or email column.", "warning")
        return render_template("lp_investors_import.html")

    import_label = request.form.get("import_label", "").strip()
    return render_template("lp_investors_import.html", preview=preview, col_map=col_map,
                           row_count=len(preview), import_label=import_label)


@app.route("/lp-investors/import/confirm", methods=["POST"])
def import_lp_investors_confirm():
    i = 0
    entries = []
    while request.form.get(f"rows[{i}][name]") is not None or request.form.get(f"rows[{i}][email]") is not None:
        entries.append({
            "name":                  request.form.get(f"rows[{i}][name]", "").strip(),
            "company":               request.form.get(f"rows[{i}][company]", "").strip(),
            "title":                 request.form.get(f"rows[{i}][title]", "").strip(),
            "email":                 request.form.get(f"rows[{i}][email]", "").strip(),
            "phone":                 request.form.get(f"rows[{i}][phone]", "").strip(),
            "location":              request.form.get(f"rows[{i}][location]", "").strip(),
            "net_worth":             request.form.get(f"rows[{i}][net_worth]", "").strip(),
            "preferred_asset_types": request.form.get(f"rows[{i}][preferred_asset_types]", "").strip(),
            "preferred_markets":     request.form.get(f"rows[{i}][preferred_markets]", "").strip(),
            "min_check":             request.form.get(f"rows[{i}][min_check]", "").strip(),
            "max_check":             request.form.get(f"rows[{i}][max_check]", "").strip(),
            "notes":                 request.form.get(f"rows[{i}][notes]", "").strip(),
        })
        i += 1

    if not entries:
        flash("No data received. Please re-upload your file.", "warning")
        return redirect(url_for("import_lp_investors"))

    import_label = request.form.get("import_label", "").strip() or "LP Import"
    valid_entries = [e for e in entries if e["name"] or e["email"]]

    conn = get_db()
    cur = conn.execute(
        "INSERT INTO import_batches (label, type, count) VALUES (?, 'lp', ?)",
        (import_label, len(valid_entries)),
    )
    batch_id = cur.lastrowid

    imported = 0
    for entry in valid_entries:
        name = entry["name"] or entry["company"] or "Unknown"
        conn.execute(
            "INSERT INTO lp_investors (name, title, company, email, phone, location, "
            "net_worth, preferred_asset_types, preferred_markets, min_check, max_check, "
            "source, status, notes, import_batch_id) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Excel Import', 'New', ?, ?)",
            (name, entry["title"], entry["company"], entry["email"], entry["phone"],
             entry["location"], entry["net_worth"], entry["preferred_asset_types"],
             entry["preferred_markets"], entry["min_check"], entry["max_check"],
             entry["notes"], batch_id),
        )
        imported += 1
    conn.commit(); conn.close()
    flash(f"Imported {imported} LP investors as '{import_label}'.", "success")
    return redirect(url_for("lp_investors"))


# ── GP Contacts CSV Export ─────────────────────────────────────────────────────

@app.route("/contacts/export.csv")
def export_gp_contacts():
    conn = get_db()
    rows = conn.execute("SELECT * FROM gp_contacts ORDER BY name").fetchall()
    conn.close()
    cols = ["id", "name", "title", "company", "email", "phone", "location", "aum_range",
            "strategy", "min_check", "max_check", "source", "status", "notes",
            "last_contacted", "next_contact_date", "created_at"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cols)
    for r in rows:
        writer.writerow([r[c] for c in cols])
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=gp_contacts.csv"})


with app.app_context():
    init_db()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="127.0.0.1", port=port, debug=os.environ.get("FLASK_DEBUG", "0") == "1")
